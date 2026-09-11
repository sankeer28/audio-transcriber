import os
import sys
import zipfile
import shutil
import re
import warnings
import json
import subprocess
import tempfile
import argparse
import io
import contextlib
from pathlib import Path
from pptx import Presentation
from faster_whisper import WhisperModel
import ctranslate2
from tqdm import tqdm

# openai-whisper (and its torch dependency) is optional - it is only needed for
# TRANSCRIPTION_ENGINE = "standard" and for the fallback path in transcribe_audio.
try:
    import whisper
except ImportError:
    whisper = None

try:
    import torch
except ImportError:
    torch = None

# Fix Windows console encoding issues
if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8')

# Suppress CUDA/Triton warnings for cleaner output
warnings.filterwarnings("ignore", category=UserWarning, module="whisper")
warnings.filterwarnings("ignore", category=UserWarning, module="faster_whisper")

# ⚙️ CONFIGURATION SETTINGS - Edit these for easy customization
# 📂 Folder Settings
PPTX_FOLDER = "presentations"   # input folder
OUTPUT_FOLDER = "output"        # output folder

# 🚀 Transcription Engine Selection
# "auto"           = pick the best available (recommended - see resolve_engine below)
# "whisper.cpp"    = GPU via Vulkan; works on AMD/Intel/NVIDIA. Needs `python setup.py`
# "faster-whisper" = CPU, or NVIDIA GPU via CUDA. Downloads its own model on first run
# "standard"       = openai-whisper (original implementation, needs torch)
TRANSCRIPTION_ENGINE = "auto"  # Options: "auto", "whisper.cpp", "faster-whisper", "standard"

# 🖥️ whisper.cpp Settings (only used when the whisper.cpp engine is active)
_EXE = ".exe" if sys.platform == "win32" else ""
WHISPERCPP_BIN = os.path.join("whispercpp", "whisper-cli" + _EXE)
WHISPERCPP_MODEL = os.path.join("models", "ggml-large-v3.bin")
WHISPERCPP_GPU_DEVICE = 0   # Vulkan device index (0 = first GPU listed at startup)
WHISPERCPP_THREADS = 0      # CPU threads for non-GPU parts (0 = auto-detect)
# Text context carried between 30s windows. This is the main cause of Whisper's
# repetition loops, where it repeats one sentence for minutes and never
# transcribes the real speech. 0 disables it and is strongly recommended:
# on a test lecture it recovered 532 words of real content instead of 200.
# Set to -1 for the library default if you need cross-window context.
WHISPERCPP_MAX_CONTEXT = 0
FFMPEG_BIN = "ffmpeg"       # ffmpeg executable, used to convert input to 16kHz mono WAV
AUTO_DOWNLOAD_MODEL = True  # Offer to fetch the GGML model when the GPU binaries are
                            # installed but the model is missing (asks first)

# 🎯 Whisper Model Settings
# "large-v3" is the most accurate model. "large-v3-turbo" is ~4x faster with a
# small accuracy cost; "medium" / "small" trade more accuracy for more speed.
WHISPER_MODEL = "large-v3"      # Options: "tiny", "base", "small", "medium", "large-v3", "large-v3-turbo"
FORCE_LANGUAGE = "en"           # Force language to prevent mixing (None for auto-detect)

# ⚡ Performance Settings
FORCE_DEVICE = "cpu"             # Options: None (auto), "cpu", "cuda" (force specific device)
USE_HALF_PRECISION = False       # fp16 for 30-50% speed boost (minimal accuracy loss)
CPU_THREADS = 0                  # CPU worker threads (0 = auto-detect from your core count)
CPU_COMPUTE_TYPE = "int8"        # "int8" (fast, near-identical quality) or "float32" (slowest, reference quality)
GPU_BEST_OF = 3                 # Decoding attempts on GPU (higher = more accurate, slower)
GPU_BEAM_SIZE = 5               # Beam search size on GPU
CPU_BEST_OF = 3                 # Decoding attempts on CPU
CPU_BEAM_SIZE = 5               # Beam search size on CPU (5 = Whisper's reference setting)

# 🎚️ Quality Settings
TEMPERATURE = 0.0               # 0.0 = deterministic, 0.1-1.0 = more creative
ENABLE_WORD_TIMESTAMPS = True   # Get word-level timing data

# 🧹 Transcript Cleanup - Whisper sometimes loops on a phrase and repeats it
# dozens of times. Each finished transcript is cleaned automatically.
AUTO_CLEAN_TRANSCRIPTS = True   # Run clean_transcripts.py on each output file
CLEAN_WRITE_BACKUP = False      # Keep a '<name>_backup.txt' of the raw transcript

# 🖥️ Command line overrides (all optional - defaults come from the settings above)
INPUT_PATH = PPTX_FOLDER

def parse_args():
    ap = argparse.ArgumentParser(
        description="Transcribe PowerPoint, MP4 and MP3 files to text.",
        epilog="Examples:\n"
               "  python main.py                        process the presentations/ folder\n"
               "  python main.py lecture.mp4            process one file\n"
               "  python main.py \"WEEK 1/videos\" -o \"WEEK 1\"\n"
               "  python main.py -i videos -o transcripts --engine faster-whisper",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", nargs="?", default=None,
                    help=f"file or folder to transcribe (default: {PPTX_FOLDER})")
    ap.add_argument("-i", "--input", dest="input_opt", default=None,
                    help="same as the positional argument")
    ap.add_argument("-o", "--output", default=None,
                    help=f"folder to write transcripts into (default: {OUTPUT_FOLDER})")
    ap.add_argument("--engine", choices=["auto", "whisper.cpp", "faster-whisper", "standard"],
                    help="override the transcription engine")
    ap.add_argument("--model", help="path to a GGML model (whisper.cpp engine)")
    ap.add_argument("--no-clean", action="store_true",
                    help="keep raw output; do not remove repeated text")
    ap.add_argument("--backup", action="store_true",
                    help="also write '<name>_backup.txt' with the uncleaned transcript")
    return ap.parse_args()

# Only parse when run directly, so importing main.py from another script is safe
if __name__ == "__main__":
    _args = parse_args()
    INPUT_PATH = _args.input_opt or _args.input or PPTX_FOLDER
    if _args.output:
        OUTPUT_FOLDER = _args.output
    if _args.engine:
        TRANSCRIPTION_ENGINE = _args.engine
    if _args.model:
        WHISPERCPP_MODEL = _args.model
    if _args.no_clean:
        AUTO_CLEAN_TRANSCRIPTS = False
    if _args.backup:
        CLEAN_WRITE_BACKUP = True

# Only auto-create the default input folder; never turn a given path into one
if INPUT_PATH == PPTX_FOLDER:
    os.makedirs(PPTX_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

def cuda_available():
    """Detect CUDA without requiring torch (ctranslate2 ships with faster-whisper)."""
    try:
        return ctranslate2.get_cuda_device_count() > 0
    except Exception:
        return torch is not None and torch.cuda.is_available()

def clear_gpu_cache():
    """Release cached GPU memory when torch is installed and CUDA is in use."""
    if device == "cuda" and torch is not None and torch.cuda.is_available():
        torch.cuda.empty_cache()

# ⚡ Load Whisper model with optimal device selection
def get_optimal_device(verbose=True):
    # Check if user forced a specific device
    if FORCE_DEVICE:
        if FORCE_DEVICE == "cuda" and not cuda_available():
            print("[!] CUDA requested but not available, falling back to CPU")
            return "cpu"
        if verbose:
            print(f"[*] Forced device: {FORCE_DEVICE}")
        return FORCE_DEVICE

    # Auto-detect best device
    if cuda_available():
        if not verbose:
            return "cuda"
        if torch is not None and torch.cuda.is_available():
            gpu_name = torch.cuda.get_device_name(0)
            gpu_memory = torch.cuda.get_device_properties(0).total_memory / 1024**3
            print(f"GPU detected: {gpu_name} ({gpu_memory:.1f}GB)")
        else:
            print(f"GPU detected ({ctranslate2.get_cuda_device_count()} CUDA device(s))")
        return "cuda"
    else:
        if verbose:
            print("No GPU detected, using CPU")
        return "cpu"

def load_standard_whisper(target_device="cpu"):
    """Load openai-whisper, with a clear error if the optional dependency is missing."""
    if whisper is None:
        raise RuntimeError(
            "openai-whisper is not installed. Install it with "
            "'pip install openai-whisper torch' to use the 'standard' engine."
        )
    return whisper.load_model(WHISPER_MODEL, device=target_device)

# ⚡ Load the appropriate Whisper model based on selected engine
model = None
faster_model = None

def fetch_whispercpp_model():
    """Download the GGML model for the GPU engine, asking first (it is GBs).

    Returns True if the model is present afterwards. Declining, a non-interactive
    shell, or a failed download all return False so the caller can fall back.
    """
    try:
        from setup import MODEL_BASE, MODELS, download
    except ImportError:
        print("[!] setup.py not found - cannot auto-download the model")
        return False

    filename = os.path.basename(WHISPERCPP_MODEL)
    size = next((s for f, s in MODELS.values() if f == filename), "several GB")

    print(f"\n[Setup] GPU binaries are installed, but the model is missing:")
    print(f"        {WHISPERCPP_MODEL}  ({size})")
    if not sys.stdin.isatty():
        print("        Run 'python setup.py' to download it. Using CPU for now.")
        return False
    try:
        answer = input(f"[Setup] Download it now? [Y/n] ").strip().lower()
    except (EOFError, KeyboardInterrupt):
        # No one there to answer (piped stdin, CI, Ctrl+C at the prompt)
        print("\n        Run 'python setup.py' to download it. Using CPU for now.")
        return False
    if answer not in ("", "y", "yes"):
        print("        Skipped. Using CPU instead.")
        return False

    os.makedirs(os.path.dirname(WHISPERCPP_MODEL) or ".", exist_ok=True)
    try:
        download(MODEL_BASE + filename, WHISPERCPP_MODEL)
    except KeyboardInterrupt:
        print("\n[!] Download cancelled. Using CPU instead.")
        return False
    except Exception as e:
        print(f"[!] Download failed: {e}")
        return False
    print(f"[OK] Saved to {WHISPERCPP_MODEL}")
    return True

def whispercpp_ready():
    """True when the GPU engine has everything it needs to run.

    If the binaries and ffmpeg are in place but only the model is missing, offer
    to download it rather than silently dropping to the much slower CPU engine.
    """
    if not os.path.isfile(WHISPERCPP_BIN) or shutil.which(FFMPEG_BIN) is None:
        return False
    if os.path.isfile(WHISPERCPP_MODEL):
        return True
    if AUTO_DOWNLOAD_MODEL:
        return fetch_whispercpp_model()
    return False

def resolve_engine(choice):
    """Turn TRANSCRIPTION_ENGINE into a concrete engine, explaining the decision.

    "auto" prefers the GPU engine when `python setup.py` has installed it, then
    falls back to faster-whisper, which downloads its own model on first run.
    """
    if choice != "auto":
        return choice
    if whispercpp_ready():
        print("[Auto] Using whisper.cpp on GPU (Vulkan)")
        return "whisper.cpp"
    if cuda_available():
        print("[Auto] Using faster-whisper on your NVIDIA GPU")
        return "faster-whisper"
    print("[Auto] Using faster-whisper on CPU. For much faster GPU transcription, "
          "run: python setup.py")
    return "faster-whisper"

TRANSCRIPTION_ENGINE = resolve_engine(TRANSCRIPTION_ENGINE)

# FORCE_DEVICE/device only apply to the faster-whisper and standard engines
device = get_optimal_device(verbose=TRANSCRIPTION_ENGINE != "whisper.cpp")

# Auto-detect a sensible thread count when left at 0
if not CPU_THREADS:
    CPU_THREADS = os.cpu_count() or 4
if not WHISPERCPP_THREADS:
    WHISPERCPP_THREADS = os.cpu_count() or 4

if TRANSCRIPTION_ENGINE == "whisper.cpp":
    # No model is loaded in-process; whisper-cli loads it per file.
    if (AUTO_DOWNLOAD_MODEL and os.path.isfile(WHISPERCPP_BIN)
            and not os.path.isfile(WHISPERCPP_MODEL)):
        fetch_whispercpp_model()
    missing = []
    if not os.path.isfile(WHISPERCPP_BIN):
        missing.append(f"binary: {WHISPERCPP_BIN}")
    if not os.path.isfile(WHISPERCPP_MODEL):
        missing.append(f"model:  {WHISPERCPP_MODEL}")
    if shutil.which(FFMPEG_BIN) is None:
        missing.append("ffmpeg on PATH")
    if missing:
        raise SystemExit("[!] whisper.cpp engine is missing:\n    "
                         + "\n    ".join(missing)
                         + "\n\n    Run 'python setup.py' to install it, or set "
                           "TRANSCRIPTION_ENGINE = \"auto\" in main.py to fall back to CPU.")
    print(f"Using whisper.cpp ({os.path.basename(WHISPERCPP_MODEL)}) on Vulkan device {WHISPERCPP_GPU_DEVICE}")
    print(f"[OK] GPU acceleration via Vulkan")
elif TRANSCRIPTION_ENGINE == "faster-whisper":
    print(f"Loading faster-whisper model ({WHISPER_MODEL}) on {device}...")
    # For faster-whisper, we need to specify compute type
    compute_type = CPU_COMPUTE_TYPE if device == "cpu" else "float16"
    faster_model = WhisperModel(WHISPER_MODEL, device=device, compute_type=compute_type,
                                cpu_threads=CPU_THREADS)
    print(f"[OK] Using faster-whisper with {compute_type} precision"
          + (f" on {CPU_THREADS} threads" if device == "cpu" and CPU_THREADS else ""))
elif TRANSCRIPTION_ENGINE == "standard":
    print(f"Loading standard openai-whisper model ({WHISPER_MODEL}) on {device}...")
    model = load_standard_whisper(device)
    print(f"[OK] Using standard openai-whisper")
else:
    print(f"[!] Invalid TRANSCRIPTION_ENGINE '{TRANSCRIPTION_ENGINE}', falling back to standard")
    TRANSCRIPTION_ENGINE = "standard"
    model = load_standard_whisper(device)

def extract_text_from_pptx(pptx_path):
    """Extract all text from slides in a PPTX."""
    prs = Presentation(pptx_path)
    texts = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    slide_texts.append(text)
        if slide_texts:
            texts.append(f"--- Slide {slide_num} ---\n" + "\n".join(slide_texts))
    return "\n\n".join(texts)

def extract_audio_from_pptx(pptx_path, temp_dir):
    """Extract embedded audio files from PPTX (wav, mp3, m4a) with proper ordering."""
    audio_files = []
    with zipfile.ZipFile(pptx_path, "r") as zip_ref:
        media_files = []
        for file in zip_ref.namelist():
            if file.startswith("ppt/media/") and file.lower().endswith((".wav", ".mp3", ".m4a")):
                media_files.append(file)

        # Sort by the numeric part in filename (media1, media2, etc.)
        def get_media_number(filename):
            match = re.search(r'media(\d+)', filename)
            return int(match.group(1)) if match else 0

        media_files.sort(key=get_media_number)

        for file in media_files:
            extracted_path = os.path.join(temp_dir, os.path.basename(file))
            with open(extracted_path, "wb") as f:
                f.write(zip_ref.read(file))
            audio_files.append(extracted_path)

    return audio_files


def get_checkpoint_file(audio_path):
    """Get checkpoint file path for a specific audio file."""
    base_name = Path(audio_path).stem
    checkpoint_file = Path(OUTPUT_FOLDER) / f"{base_name}_checkpoint.json"
    return checkpoint_file

def load_checkpoint(checkpoint_file):
    """Load existing checkpoint if available."""
    if checkpoint_file.exists():
        with open(checkpoint_file, 'r', encoding='utf-8') as f:
            return json.load(f)
    return None

def save_checkpoint(checkpoint_file, checkpoint_data):
    """Save checkpoint data to file."""
    with open(checkpoint_file, 'w', encoding='utf-8') as f:
        json.dump(checkpoint_data, f, ensure_ascii=False, indent=2)

def clean_output_file(output_path):
    """Strip Whisper's repetition loops from a finished transcript, in place."""
    if not AUTO_CLEAN_TRANSCRIPTS:
        return
    try:
        import clean_transcripts
    except ImportError:
        print("[!] clean_transcripts.py not found - skipping cleanup")
        return

    try:
        before = os.path.getsize(output_path)
        # The individual passes log every removal; keep that out of the run log
        with contextlib.redirect_stdout(io.StringIO()):
            changed = clean_transcripts.clean_transcript_file(
                output_path, write_backup=CLEAN_WRITE_BACKUP, quiet=True
            )
        if changed:
            after = os.path.getsize(output_path)
            saved = (before - after) / before * 100 if before else 0
            print(f"[Cleaned] Removed repeated text ({saved:.1f}% smaller)")
    except Exception as e:
        # Never lose a transcript because cleanup failed
        print(f"[!] Cleanup failed for {output_path}: {e}")

def convert_to_wav16k(src_path, dst_path):
    """Decode any input (mp3/mp4/m4a/wav) to the 16kHz mono WAV whisper.cpp expects."""
    subprocess.run(
        [FFMPEG_BIN, "-y", "-loglevel", "error", "-i", src_path,
         "-ar", "16000", "-ac", "1", "-c:a", "pcm_s16le", dst_path],
        check=True
    )

def transcribe_with_whispercpp(audio_path, checkpoint_file):
    """Transcribe via whisper.cpp on the GPU (Vulkan), with checkpoint/resume support."""
    checkpoint = load_checkpoint(checkpoint_file)
    if checkpoint and checkpoint.get("segments"):
        resume_from = checkpoint["segments"][-1]["end"]
        print(f"[Checkpoint] Found existing progress! "
              f"{len(checkpoint['segments'])} segments ({resume_from:.1f}s)")
    else:
        checkpoint = {"segments": [], "metadata": {}}
        resume_from = 0.0
        print(f"[Transcribing] Starting new transcription...")

    temp_dir = tempfile.mkdtemp(prefix="wcpp_")
    try:
        wav_path = os.path.join(temp_dir, "audio.wav")
        convert_to_wav16k(audio_path, wav_path)
        out_base = os.path.join(temp_dir, "result")

        cmd = [
            WHISPERCPP_BIN,
            "-m", WHISPERCPP_MODEL,
            "-f", wav_path,
            "-t", str(WHISPERCPP_THREADS),
            "-bs", str(GPU_BEAM_SIZE),
            "-dev", str(WHISPERCPP_GPU_DEVICE),
            "-mc", str(WHISPERCPP_MAX_CONTEXT),
            "-oj", "-of", out_base,
            "-pp",
        ]
        if FORCE_LANGUAGE:
            cmd += ["-l", FORCE_LANGUAGE]
        # Resume: skip audio we already transcribed. whisper.cpp reports timestamps
        # relative to this offset, so it is added back when parsing below.
        if resume_from > 0:
            cmd += ["-ot", str(int(resume_from * 1000))]

        proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE,
                                text=True, encoding="utf-8", errors="replace")
        progress_re = re.compile(r"progress\s*=\s*(\d+)%")
        with tqdm(total=100, desc="Transcribing (GPU)", unit="%",
                  bar_format="{l_bar}{bar}| {n:.0f}/100% [{elapsed}<{remaining}]") as pbar:
            for line in proc.stderr:
                m = progress_re.search(line)
                if m:
                    pct = int(m.group(1))
                    pbar.update(max(0, pct - pbar.n))
            proc.wait()
            if proc.returncode == 0:
                pbar.update(max(0, 100 - pbar.n))

        if proc.returncode != 0:
            raise RuntimeError(f"whisper-cli failed with exit code {proc.returncode}")

        with open(out_base + ".json", "r", encoding="utf-8") as f:
            data = json.load(f)

        for item in data.get("transcription", []):
            offsets = item.get("offsets", {})
            checkpoint["segments"].append({
                "start": resume_from + offsets.get("from", 0) / 1000.0,
                "end": resume_from + offsets.get("to", 0) / 1000.0,
                "text": item.get("text", ""),
            })

        checkpoint["metadata"] = {"file": str(audio_path), "engine": "whisper.cpp"}
        save_checkpoint(checkpoint_file, checkpoint)
        print(f"[Completed] Transcription finished! Total segments: {len(checkpoint['segments'])}")

        text = " ".join(seg["text"] for seg in checkpoint["segments"])
        print(f"[Cleanup] Removing checkpoint file...")
        checkpoint_file.unlink(missing_ok=True)
        return {"text": text.strip()}
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)

def transcribe_single_file(audio_path):
    """Transcribe a single audio/video file using the selected engine with checkpoint support."""
    checkpoint_file = get_checkpoint_file(audio_path)

    if TRANSCRIPTION_ENGINE == "whisper.cpp":
        return transcribe_with_whispercpp(audio_path, checkpoint_file)
    elif TRANSCRIPTION_ENGINE == "faster-whisper":
        # Check for existing checkpoint
        checkpoint = load_checkpoint(checkpoint_file)
        if checkpoint and checkpoint.get("segments"):
            resume_from = checkpoint["segments"][-1]["end"]
            print(f"[Checkpoint] Found existing progress! "
                  f"{len(checkpoint['segments'])} segments ({resume_from:.1f}s)")
        else:
            checkpoint = {"segments": [], "metadata": {}}
            resume_from = 0.0
            print(f"[Transcribing] Starting new transcription...")

        transcribe_kwargs = dict(
            language=FORCE_LANGUAGE,
            task="transcribe",
            temperature=TEMPERATURE,
            beam_size=GPU_BEAM_SIZE if device == "cuda" else CPU_BEAM_SIZE,
            word_timestamps=ENABLE_WORD_TIMESTAMPS
        )

        # Resume by seeking the decoder to where we left off, so already-transcribed
        # audio is not decoded a second time. clip_timestamps with an odd number of
        # values runs from that offset to the end of the file, and the segment
        # timestamps it yields are absolute.
        skip_until = 0.0
        if resume_from > 0:
            try:
                segments, info = faster_model.transcribe(
                    audio_path, clip_timestamps=[resume_from], **transcribe_kwargs
                )
            except TypeError:
                # Older faster-whisper without clip_timestamps support: re-decode from
                # the start and drop what we already have.
                print("[!] This faster-whisper build cannot seek; re-decoding from start")
                skip_until = resume_from
                segments, info = faster_model.transcribe(audio_path, **transcribe_kwargs)
        else:
            segments, info = faster_model.transcribe(audio_path, **transcribe_kwargs)

        # Store metadata
        checkpoint["metadata"] = {
            "duration": info.duration,
            "language": info.language,
            "file": str(audio_path)
        }

        print(f"[Info] Duration: {info.duration:.1f}s | Language: {info.language}")
        if resume_from == 0:
            print(f"[Checkpoint] Saving to: {checkpoint_file}")

        # Create progress bar
        with tqdm(total=int(info.duration), desc="Transcribing", unit="s",
                  bar_format="{l_bar}{bar}| {n:.0f}/{total:.0f}s [{elapsed}<{remaining}]",
                  initial=int(resume_from)) as pbar:

            new_segments = 0
            last_position = int(resume_from)

            for segment in segments:
                # Only relevant on the no-seek fallback path
                if segment.end <= skip_until:
                    continue

                # Add new segment to checkpoint
                segment_data = {
                    'start': segment.start,
                    'end': segment.end,
                    'text': segment.text
                }
                checkpoint["segments"].append(segment_data)
                new_segments += 1

                # Save checkpoint every 10 segments (balance between safety and I/O)
                if new_segments % 10 == 0:
                    save_checkpoint(checkpoint_file, checkpoint)

                # Update progress bar
                current_position = int(segment.end)
                pbar.update(max(0, current_position - last_position))
                last_position = current_position

            # Final save
            save_checkpoint(checkpoint_file, checkpoint)

        print(f"[Completed] Transcription finished! Total segments: {len(checkpoint['segments'])}")
        
        # Extract text from checkpoint
        text = " ".join([seg["text"] for seg in checkpoint["segments"]])
        
        # Clean up checkpoint on successful completion
        print(f"[Cleanup] Removing checkpoint file...")
        checkpoint_file.unlink(missing_ok=True)
        
        return {"text": text.strip()}
    else:
        # Standard openai-whisper
        if device == "cuda":
            result = model.transcribe(
                audio_path,
                language=FORCE_LANGUAGE,
                task="transcribe",
                temperature=TEMPERATURE,
                best_of=GPU_BEST_OF,
                beam_size=GPU_BEAM_SIZE,
                word_timestamps=ENABLE_WORD_TIMESTAMPS,
                fp16=USE_HALF_PRECISION
            )
        else:
            result = model.transcribe(
                audio_path,
                language=FORCE_LANGUAGE,
                task="transcribe",
                temperature=TEMPERATURE,
                best_of=CPU_BEST_OF,
                beam_size=CPU_BEAM_SIZE,
                word_timestamps=ENABLE_WORD_TIMESTAMPS,
                fp16=False
            )
        return result

def transcribe_audio(audio_files):
    """Run Whisper transcription with hybrid CPU/GPU optimization and progress bar."""
    transcripts = []

    # Create progress bar
    progress_bar = tqdm(
        audio_files,
        desc="[Audio] Transcribing",
        unit="file",
        bar_format="{l_bar}{bar}| {n_fmt}/{total_fmt} [{elapsed}<{remaining}]"
    )

    for audio_path in progress_bar:
        # Update progress bar description with current file
        filename = os.path.basename(audio_path)
        media_match = re.search(r'media(\d+)', filename)
        media_num = media_match.group(1) if media_match else "unknown"

        progress_bar.set_description(f"[Audio] Processing {media_num}")

        try:
            # Use the unified transcribe function
            result = transcribe_single_file(audio_path)

        except KeyboardInterrupt:
            raise
        except Exception as e:
            progress_bar.write(f"[!] Error with {filename}: {e}")
            result = None
            # Fallback - try with standard whisper if the GPU/faster path fails
            if TRANSCRIPTION_ENGINE in ("faster-whisper", "whisper.cpp"):
                progress_bar.write(f"[!] Falling back to standard whisper for {filename}...")
                try:
                    cpu_model = load_standard_whisper("cpu")
                    result = cpu_model.transcribe(
                        audio_path,
                        language=FORCE_LANGUAGE,
                        temperature=TEMPERATURE
                    )
                except KeyboardInterrupt:
                    raise
                except Exception as fallback_error:
                    progress_bar.write(f"[!] Fallback also failed for {filename}: {fallback_error}")

            # Skip this file rather than aborting the whole run
            if result is None:
                progress_bar.write(f"[!] Skipping {filename}")
                transcripts.append(
                    f"--- Audio {media_num} Transcript ---\n[Transcription failed: {e}]"
                )
                clear_gpu_cache()
                continue

        transcripts.append(f"--- Audio {media_num} Transcript ---\n{result['text'].strip()}")

        # Clear GPU cache periodically if using CUDA
        clear_gpu_cache()

    progress_bar.close()
    return "\n\n".join(transcripts)

def process_pptx(pptx_path):
    """Process one PowerPoint: text + audio → output TXT file."""
    print(f"\n[PPTX] Processing {pptx_path}...")
    base_name = os.path.splitext(os.path.basename(pptx_path))[0]
    temp_dir = os.path.join(OUTPUT_FOLDER, base_name + "_media")
    os.makedirs(temp_dir, exist_ok=True)

    try:
        # Slide text
        text_content = extract_text_from_pptx(pptx_path)

        # Embedded audio
        audio_files = extract_audio_from_pptx(pptx_path, temp_dir)
        transcript = ""
        if audio_files:
            try:
                transcript = transcribe_audio(audio_files)
            except KeyboardInterrupt:
                raise
            except Exception as e:
                print(f"[ERROR] Audio transcription failed for {pptx_path}: {e}")
                print("[!] Saving slide text only")

        # Combine output with improved organization
        final_output = []
        if text_content:
            final_output.append("### PowerPoint Slide Content (In Order) ###\n" + text_content)
        if transcript:
            final_output.append("\n### Audio Transcripts (In Chronological Order) ###\n" + transcript)

        # Add summary note about ordering
        if text_content and transcript:
            final_output.append("\n### Note ###\nAudio transcripts are generated by OpenAI Whisper.")

        # Save result
        output_path = os.path.join(OUTPUT_FOLDER, base_name + ".txt")
        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n\n".join(final_output))

        print(f"[OK] Saved results to {output_path}")
        clean_output_file(output_path)

    finally:
        # Cleanup extracted media
        shutil.rmtree(temp_dir, ignore_errors=True)

def process_mp3(mp3_path):
    """Process standalone MP3 file: transcribe audio → output TXT file."""
    print(f"\n[MP3] Processing {mp3_path}...")
    base_name = os.path.splitext(os.path.basename(mp3_path))[0]

    try:
        # Transcribe the MP3 file
        print("[Audio] Transcribing...")
        result = transcribe_single_file(mp3_path)

        # Save transcription
        output_path = os.path.join(OUTPUT_FOLDER, base_name + ".txt")
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(f"### MP3 Audio Transcription ###\n\n{result['text'].strip()}")

        print(f"[OK] Saved transcription to {output_path}")
        clean_output_file(output_path)

        # Clear GPU cache if using CUDA
        clear_gpu_cache()

    except Exception as e:
        print(f"[ERROR] Error processing {mp3_path}: {e}")

def process_mp4(mp4_path):
    """Process MP4 video file: extract and transcribe audio → output TXT file."""
    print(f"\n[MP4] Processing {mp4_path}...")
    base_name = os.path.splitext(os.path.basename(mp4_path))[0]

    try:
        # Transcribe the MP4 file (Whisper extracts audio automatically)
        print("[Audio] Extracting and transcribing...")
        result = transcribe_single_file(mp4_path)

        # Save transcription
        output_path = os.path.join(OUTPUT_FOLDER, base_name + ".txt")
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(result['text'].strip())

        print(f"[OK] Saved transcription to {output_path}")
        clean_output_file(output_path)

        # Clear GPU cache if using CUDA
        clear_gpu_cache()

    except Exception as e:
        print(f"[ERROR] Error processing {mp4_path}: {e}")

HANDLERS = {".pptx": process_pptx, ".mp3": process_mp3, ".mp4": process_mp4}

def process_one(path):
    """Dispatch a single input file to the right handler by extension."""
    handler = HANDLERS.get(os.path.splitext(path)[1].lower())
    if handler is None:
        print(f"[!] Unsupported file type: {path}")
        return False
    handler(path)
    return True

def main():
    # A single file was given - just do that one
    if os.path.isfile(INPUT_PATH):
        print(f"[Files] Processing 1 file -> {OUTPUT_FOLDER}")
        process_one(INPUT_PATH)
        return

    if not os.path.isdir(INPUT_PATH):
        print(f"[!] Input not found: {INPUT_PATH}")
        return

    # Scan the folder for all supported file types
    entries = sorted(os.listdir(INPUT_PATH))
    by_type = {ext: [f for f in entries if f.lower().endswith(ext)]
               for ext in HANDLERS}

    total_files = sum(len(v) for v in by_type.values())
    if total_files == 0:
        print(f"[!] No .pptx, .mp3, or .mp4 files found in {INPUT_PATH}")
        return

    print(f"[Files] Found {len(by_type['.pptx'])} PPTX, {len(by_type['.mp3'])} MP3, "
          f"{len(by_type['.mp4'])} MP4 files in {INPUT_PATH} -> {OUTPUT_FOLDER}")

    for ext in (".pptx", ".mp3", ".mp4"):
        for file in by_type[ext]:
            process_one(os.path.join(INPUT_PATH, file))

if __name__ == "__main__":
    main()
